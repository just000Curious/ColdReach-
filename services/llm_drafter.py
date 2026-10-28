import json
import os
import re
import time
import base64
import requests
from typing import List, Dict, Any
from dotenv import load_dotenv

load_dotenv()

from services.db import add_token_usage

# Anchor to project root for additional project details
_DIR = os.path.dirname(os.path.abspath(__file__))
_PROJECT_DETAILS_FILE = os.path.join(os.path.dirname(_DIR), "project_resume_details.md")

def _load_project_details() -> str:
    """Load additional project details (e.g. 7-agent LangGraph RAG platform) if available."""
    if os.path.exists(_PROJECT_DETAILS_FILE):
        try:
            with open(_PROJECT_DETAILS_FILE, "r", encoding="utf-8") as f:
                content = f.read()
                return content[:4000]
        except Exception:
            pass
    return ""



class EmailDraftResponse:
    """Container for the LLM's structured email draft."""
    def __init__(self, hiring_manager_name, company_name, subject, body):
        self.hiring_manager_name = hiring_manager_name
        self.company_name = company_name
        self.subject = subject
        self.body = body


class ExtractedInfo:
    """Container for auto-extracted job info from JD or poster."""
    def __init__(self, email="", job_title="", company_name="", job_description=""):
        self.email = email
        self.job_title = job_title
        self.company_name = company_name
        self.job_description = job_description


def _strip_markdown_fences(text: str) -> str:
    """Gemini sometimes wraps JSON in ```json ... ``` fences. Strip them."""
    text = text.strip()
    if text.startswith("```"):
        text = re.sub(r"^```\w*\s*\n?", "", text)
        text = re.sub(r"\n?```\s*$", "", text)
    return text.strip()


def _get_api_key() -> str:
    api_key = os.environ.get("GEMINI_API_KEY", "").strip('" ')
    if not api_key:
        raise ValueError("GEMINI_API_KEY is missing. Please set it in your .env file.")
    return api_key


def _get_model_name() -> str:
    return os.environ.get("GEMINI_MODEL", "gemini-3.5-flash").strip('" ')


def call_gemini(prompt: str, json_mode: bool = False, temperature: float = 0.2) -> str:
    """Call Gemini text API with automatic retry on transient errors."""
    api_key = _get_api_key()
    model = _get_model_name()
    url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent?key={api_key}"

    payload = {
        "contents": [{"parts": [{"text": prompt}]}],
        "generationConfig": {"temperature": temperature},
    }
    if json_mode:
        payload["generationConfig"]["responseMimeType"] = "application/json"

    headers = {"Content-Type": "application/json"}

    last_error = None
    for attempt in range(3):
        try:
            response = requests.post(url, json=payload, headers=headers, timeout=60)
            if response.status_code in (503, 429, 500):
                time.sleep(2 ** (attempt + 1))
                last_error = f"{response.status_code}: {response.text[:200]}"
                continue
            response.raise_for_status()
            data = response.json()
            if "usageMetadata" in data:
                add_token_usage(data["usageMetadata"].get("totalTokenCount", 0))
            raw_text = data["candidates"][0]["content"]["parts"][0]["text"]
            return _strip_markdown_fences(raw_text)
        except requests.exceptions.Timeout:
            last_error = "Timeout"
            if attempt < 2:
                time.sleep(2 ** (attempt + 1))
                continue
        except requests.exceptions.HTTPError as e:
            last_error = str(e)
            if attempt < 2:
                time.sleep(2 ** (attempt + 1))
                continue
            raise
    raise ValueError(f"Gemini API failed after 3 retries. Last error: {last_error}")


def call_gemini_vision(prompt: str, image_bytes: bytes, json_mode: bool = False) -> str:
    """Call Gemini with an image (multimodal) for poster/flyer analysis."""
    api_key = _get_api_key()
    model = _get_model_name()
    url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent?key={api_key}"

    b64_image = base64.b64encode(image_bytes).decode("utf-8")

    payload = {
        "contents": [{
            "parts": [
                {"text": prompt},
                {"inline_data": {"mime_type": "image/png", "data": b64_image}},
            ]
        }],
        "generationConfig": {"temperature": 0.1},
    }
    if json_mode:
        payload["generationConfig"]["responseMimeType"] = "application/json"

    headers = {"Content-Type": "application/json"}

    last_error = None
    for attempt in range(3):
        try:
            response = requests.post(url, json=payload, headers=headers, timeout=90)
            if response.status_code in (503, 429, 500):
                time.sleep(2 ** (attempt + 1))
                last_error = f"{response.status_code}: {response.text[:200]}"
                continue
            response.raise_for_status()
            data = response.json()
            if "usageMetadata" in data:
                add_token_usage(data["usageMetadata"].get("totalTokenCount", 0))
            raw_text = data["candidates"][0]["content"]["parts"][0]["text"]
            return _strip_markdown_fences(raw_text)
        except requests.exceptions.Timeout:
            last_error = "Timeout"
            if attempt < 2:
                time.sleep(2 ** (attempt + 1))
                continue
        except requests.exceptions.HTTPError as e:
            last_error = str(e)
            if attempt < 2:
                time.sleep(2 ** (attempt + 1))
                continue
            raise
    raise ValueError(f"Gemini Vision API failed after 3 retries. Last error: {last_error}")

def call_gemini_pdf(prompt: str, pdf_bytes: bytes, json_mode: bool = False) -> str:
    """Call Gemini with a PDF for robust text extraction."""
    api_key = _get_api_key()
    model = _get_model_name()
    url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent?key={api_key}"

    b64_pdf = base64.b64encode(pdf_bytes).decode("utf-8")

    payload = {
        "contents": [{
            "parts": [
                {"text": prompt},
                {"inline_data": {"mime_type": "application/pdf", "data": b64_pdf}},
            ]
        }],
        "generationConfig": {"temperature": 0.1},
    }
    if json_mode:
        payload["generationConfig"]["responseMimeType"] = "application/json"

    headers = {"Content-Type": "application/json"}

    last_error = None
    for attempt in range(3):
        try:
            response = requests.post(url, json=payload, headers=headers, timeout=120)
            if response.status_code in (503, 429, 500):
                time.sleep(2 ** (attempt + 1))
                last_error = f"{response.status_code}: {response.text[:200]}"
                continue
            response.raise_for_status()
            data = response.json()
            if "usageMetadata" in data:
                add_token_usage(data["usageMetadata"].get("totalTokenCount", 0))
            raw_text = data["candidates"][0]["content"]["parts"][0]["text"]
            return _strip_markdown_fences(raw_text)
        except requests.exceptions.Timeout:
            last_error = "Timeout"
            if attempt < 2:
                time.sleep(2 ** (attempt + 1))
                continue
        except requests.exceptions.HTTPError as e:
            last_error = str(e)
            if attempt < 2:
                time.sleep(2 ** (attempt + 1))
                continue
            raise
    raise ValueError(f"Gemini API failed after 3 retries. Last error: {last_error}")


# ══════════════════════════════════════════════════════════════════
# EXTRACTION FUNCTIONS (Auto-fill from JD or Poster)
# ══════════════════════════════════════════════════════════════════

def extract_from_jd(jd_text: str) -> ExtractedInfo:
    """Use Gemini to auto-extract email, job title, company from a JD."""
    prompt = f"""Analyze this job description and extract the following information.

Return ONLY a JSON object with these keys:
- "email": the HR/recruiter email address if present, otherwise empty string ""
- "job_title": the exact job title being advertised
- "company_name": the company name
- "job_description": a 2-3 sentence summary of what the job requires

If any field cannot be found, use an empty string "".

JOB DESCRIPTION:
{jd_text[:4000]}"""

    result = call_gemini(prompt, json_mode=True, temperature=0.1)
    data = json.loads(result)
    return ExtractedInfo(
        email=data.get("email", ""),
        job_title=data.get("job_title", ""),
        company_name=data.get("company_name", ""),
        job_description=data.get("job_description", ""),
    )


def extract_from_poster(image_bytes: bytes) -> ExtractedInfo:
    """Use Gemini Vision to scan a poster/flyer and extract job details."""
    prompt = """Analyze this job posting image/flyer/advertisement carefully.

Extract ALL of the following information from the image:
- "email": any email address visible in the image (look carefully for HR/recruiter emails)
- "job_title": the job title or position being advertised
- "company_name": the company or organization name
- "job_description": summarize what the job requires in 2-3 sentences

Return ONLY a JSON object with those 4 keys. If any field cannot be found, use an empty string ""."""

    result = call_gemini_vision(prompt, image_bytes, json_mode=True)
    data = json.loads(result)
    return ExtractedInfo(
        email=data.get("email", ""),
        job_title=data.get("job_title", ""),
        company_name=data.get("company_name", ""),
        job_description=data.get("job_description", ""),
    )


def extract_leads(file_bytes: bytes, filename: str) -> List[Dict[str, str]]:
    """Extract a list of roles, companies, HR emails, and skills from a raw PDF or Word file using Gemini."""
    prompt = """Analyze the attached document, which contains a list of jobs, roles, or contacts.
Extract all the leads you can find. For each lead, identify:
- "role": The job title or position.
- "company": The name of the company.
- "hr_email": The email address of the HR person, recruiter, or contact provided. If not found, leave as empty string.
- "skills": A comma-separated list of key skills required. If not found, leave as empty string.

Return ONLY a JSON array of objects, e.g. [{"role": "...", "company": "...", "hr_email": "...", "skills": "..."}]. If none found, return [].
"""

    try:
        if filename.endswith(".doc") or filename.endswith(".docx"):
            import io
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            text = "\n".join([p.text for p in doc.paragraphs])
            prompt_with_text = prompt + f"\n\nDOCUMENT CONTENT:\n{text[:20000]}"
            result = call_gemini(prompt_with_text, json_mode=True, temperature=0.1)
        else:
            result = call_gemini_pdf(prompt, file_bytes, json_mode=True)
            
        data = json.loads(result)
        if isinstance(data, list):
            return data
        elif isinstance(data, dict) and "leads" in data:
            return data["leads"]
        return []
    except Exception as e:
        print(f"Error extracting leads from document via Gemini: {e}")
        return []

# ══════════════════════════════════════════════════════════════════
# EMAIL GENERATION FUNCTIONS
# ══════════════════════════════════════════════════════════════════

def generate_cold_email(
    resume_text: str, company_info: str, target_role: str,
    sender_name: str, github_url: str, portfolio_url: str, linkedin_url: str = "", tone: str = "Professional and direct",
) -> EmailDraftResponse:
    """Generate a cold outreach email (no JD, just reaching out)."""
    sender_name = sender_name.strip() if sender_name else "[Your Name]"
    links = []
    if portfolio_url and portfolio_url.strip(): links.append(f"Portfolio: {portfolio_url.strip()}")
    if github_url and github_url.strip(): links.append(f"GitHub: {github_url.strip()}")
    if linkedin_url and linkedin_url.strip(): links.append(f"LinkedIn: {linkedin_url.strip()}")
    links_str = " | ".join(links)
    project_details = _load_project_details()

    prompt = f"""You are an expert career outreach assistant. Write a high-converting cold outreach email to an HR manager or recruiter.

RULES:
1. NEVER invent experience, projects, or degrees not in the resume or project details.
2. Keep the email punchy, professional, and under 120 words.
3. The sender's name is: {sender_name}
4. Target Role: {target_role if target_role else "AI / Machine Learning / GenAI Engineer"}
5. Company context: {company_info if company_info else "No specific company info provided"}
6. Tone/Style: {tone}

7. MANDATORY PROJECT SELECTION & ALIGNMENT:
   Analyze the target role ('{target_role}') and company context:
   - FOR GENAI / LLM / AGENTIC AI / AI ROLES:
     (Keywords: GenAI, Generative AI, LLM, Agentic, Multi-Agent, LangGraph, LangChain, RAG, Pinecone, Vector DB, Prompt Engineering, NLP, AI Engineer, AI Developer, Artificial Intelligence)
     YOU MUST EXPLICITLY HIGHLIGHT the candidate's flagship GenAI project:
     --> **AI Academic Project Mentor** (from Infosys Springboard internship): Built a 7-agent sequential LangGraph state machine pipeline (feasibility check, scope, tech stack, Gantt milestone planning, risk matrix, documentation) with RAG-based mentoring using Pinecone Serverless vector search, persistent memory in Supabase PostgreSQL, and FastAPI backend.
     --> Can also highlight his Agentic AI & Generative AI internship at Innomatics Research Labs (RAG workflows, embeddings, prompt engineering, LLM orchestration).
     --> CRITICAL: DO NOT lead with or focus on classical tabular ML projects (such as PASE energy forecasting or credit risk modeling) for GenAI/LLM roles! GenAI recruiters specifically look for agentic and RAG architecture experience.
   - FOR CLASSICAL ML / DATA SCIENCE / MLOPS ROLES:
     (Keywords: Data Scientist, Machine Learning Engineer, MLOps, Predictive Modeling, Time Series, Tabular Data)
     --> Highlight **PASE** (Predictive Analytics for Sustainable Energy: time-series wind/energy forecasting with XGBoost & Random Forest, Dockerized FastAPI) and/or **Credit Risk Modelling System** (XGBoost, Optuna tuning across 50+ trials, SMOTE, MLflow tracking).
   - FOR HYBRID AI/ML ROLES:
     --> Mention BOTH the 7-agent LangGraph GenAI pipeline and production ML models.
   - Always reference specific project names, exact technologies (LangGraph, Pinecone, FastAPI, Supabase, Docker, MLflow, XGBoost), and concrete outcomes.

EMAIL FORMAT:
Start with "Dear [HR Name / Hiring Manager],"
Write the body paragraphs highlighting the exact relevant project and skills.
CRITICAL: DO NOT include ANY sign-off, signature, "Best regards", or your name at the end. The system will append it automatically. End the email body with the final sentence's punctuation.

Return ONLY a JSON object:
- "hiring_manager_name": "Hiring Manager"
- "company_name": company name or "the company"
- "subject": compelling subject line tailored to the target role
- "body": the full email body text. VERY IMPORTANT: DO NOT use HTML tags like <br> or <p>. Use literal newline characters (\n) for line breaks.

RESUME & BACKGROUND:
{resume_text[:12000]}
{f"\nADDITIONAL PROJECT DETAILS:\n{project_details}" if project_details else ""}"""

    result = call_gemini(prompt, json_mode=True, temperature=0.3)
    data = json.loads(result)
    
    # Programmatically append the signature
    body = data.get("body", "").strip()
    signature = f"\n\nBest regards,\n{sender_name}"
    if links_str:
        signature += f"\n{links_str}"
    
    data["body"] = body + signature
    return EmailDraftResponse(**data)


def generate_jd_email(
    resume_text: str, jd_text: str, job_title: str, company_name: str,
    sender_name: str, github_url: str, portfolio_url: str, linkedin_url: str = "", tone: str = "Professional and direct",
) -> EmailDraftResponse:
    """Generate an application email based on a specific job description."""
    sender_name = sender_name.strip() if sender_name else "[Your Name]"
    links = []
    if portfolio_url and portfolio_url.strip(): links.append(f"Portfolio: {portfolio_url.strip()}")
    if github_url and github_url.strip(): links.append(f"GitHub: {github_url.strip()}")
    if linkedin_url and linkedin_url.strip(): links.append(f"LinkedIn: {linkedin_url.strip()}")
    links_str = " | ".join(links)
    project_details = _load_project_details()

    prompt = f"""You are an expert career outreach assistant. Write a high-converting job application email tailored precisely to the job description.

RULES:
1. NEVER invent experience, projects, or degrees not in the resume or project details.
2. Keep the email under 130 words. Be specific about why this candidate is a direct fit for THIS role.
3. Sender name: {sender_name}
4. Job Title: {job_title}
5. Company: {company_name}
6. Tone/Style: {tone}

7. MANDATORY PROJECT SELECTION & ALIGNMENT:
   Analyze the Job Title ('{job_title}') and Job Description text:
   - IF the JD or Job Title involves Generative AI, GenAI, LLMs, Agentic AI, Agents, LangGraph, LangChain, RAG, Vector Databases, Prompt Engineering, NLP, or AI Engineering:
     YOU MUST EXPLICITLY HIGHLIGHT the candidate's flagship GenAI project:
     --> **AI Academic Project Mentor** (from Infosys Springboard internship): Built a 7-agent sequential LangGraph state machine pipeline (feasibility check, scope, tech stack, Gantt milestones, risk matrix, documentation) with RAG-based mentoring using Pinecone Serverless vector search, persistent memory in Supabase PostgreSQL, and FastAPI backend.
     --> Can also highlight his Agentic AI & Generative AI internship at Innomatics Research Labs (RAG workflows, embeddings, prompt engineering, LLM orchestration).
     --> CRITICAL: DO NOT default to or lead with classical tabular ML projects (such as PASE energy forecasting or credit risk modeling) when applying for GenAI/LLM roles! Recruiters want to see multi-agent pipelines and RAG architecture.
   - IF the role is focused on Classical ML, Data Science, or MLOps (e.g. Data Scientist, ML Engineer, Time-Series Forecasting, Tabular Data, MLOps):
     --> Highlight **PASE** (Predictive Analytics for Sustainable Energy: time-series wind/energy forecasting with XGBoost & Random Forest, Dockerized FastAPI) and/or **Credit Risk Modelling System** (XGBoost, Optuna tuning, SMOTE, MLflow tracking).
   - IF the role is Hybrid (AI/ML Engineer, Data Scientist with GenAI):
     --> Feature BOTH the 7-agent LangGraph GenAI pipeline and production ML models.
   - Directly mirror 2-3 specific requirements from the JD with the candidate's verified stack.

EMAIL FORMAT:
Start with "Dear [HR Name / Hiring Manager],"
Write the body paragraphs.
CRITICAL: DO NOT include ANY sign-off, signature, "Best regards", or your name at the end. The system will append it automatically. End the email body with the final sentence's punctuation.

Return ONLY a JSON object:
- "hiring_manager_name": HR name if known, else "Hiring Manager"
- "company_name": "{company_name}"
- "subject": compelling subject line mentioning the role and candidate name
- "body": the full email body text. VERY IMPORTANT: DO NOT use HTML tags like <br> or <p>. Use literal newline characters (\n) for line breaks.

RESUME & BACKGROUND:
{resume_text[:12000]}

JOB DESCRIPTION:
{jd_text[:5000]}
{f"\nADDITIONAL PROJECT DETAILS:\n{project_details}" if project_details else ""}"""

    result = call_gemini(prompt, json_mode=True, temperature=0.2)
    data = json.loads(result)
    
    # Programmatically append the signature
    body = data.get("body", "").strip()
    signature = f"\n\nBest regards,\n{sender_name}"
    if links_str:
        signature += f"\n{links_str}"
    
    data["body"] = body + signature
    return EmailDraftResponse(**data)


def rewrite_email(original_body: str, app_type: str, company: str, resume_text: str, new_angle: str, sender_name: str = "", github_url: str = "", portfolio_url: str = "", linkedin_url: str = "") -> str:
    """Rewrite email with a new angle or phrasing."""
    sender_name = sender_name.strip() if sender_name else "[Your Name]"
    links = []
    if portfolio_url and portfolio_url.strip(): links.append(f"Portfolio: {portfolio_url.strip()}")
    if github_url and github_url.strip(): links.append(f"GitHub: {github_url.strip()}")
    if linkedin_url and linkedin_url.strip(): links.append(f"LinkedIn: {linkedin_url.strip()}")
    links_str = " | ".join(links)
    project_details = _load_project_details()
    
    if new_angle:
        angle_instruction = f"The user requested this specific new angle/focus for the rewrite: '{new_angle}'. Ensure this is central to the new draft."
    else:
        angle_instruction = "Try a completely new angle or emphasize different relevant skills from the resume to make it fresh and engaging."

    project_guideline = """
PROJECT SELECTION GUIDELINE:
- If the original email, company, or requested angle involves Generative AI, LLMs, Agents, LangGraph, RAG, or AI Engineering:
  Explicitly highlight the candidate's flagship GenAI project: the **AI Academic Project Mentor** (7-agent LangGraph pipeline, Pinecone Serverless RAG, Supabase PostgreSQL, FastAPI) and/or his Agentic AI internship at Innomatics Research Labs. Do not dilute with unrelated ML projects unless specifically requested.
- If the focus is classical ML or Data Science: Highlight PASE (XGBoost/Random Forest energy forecasting) or Credit Risk Modelling (XGBoost, Optuna, MLflow).
- Keep the email concise (under 120 words), complete, and end cleanly with a clear call to action."""

    if app_type == "Cold Email":
        prompt = f"""You are rewriting a Cold Outreach email to {company}. 
{angle_instruction}
{project_guideline}

Do NOT just swap synonyms. Construct a completely fresh, coherent, and polished cold outreach email.
Keep it professional, highly personalized, and under 120 words.
DO NOT use HTML tags like <br>. Use literal newline characters for spacing.
CRITICAL: DO NOT leave any sentence incomplete or cut off. End cleanly with a question or call to action.
CRITICAL: DO NOT include ANY sign-off, signature, "Best regards", your name, or any links/URLs at the end. The system will append the signature automatically. End the email body strictly with the final sentence's punctuation.

Resume Context:
{resume_text[:12000]}
{f"\nProject Details:\n{project_details}" if project_details else ""}

Original Email (MAY contain a signature, ignore it):
{original_body}

Return ONLY the rewritten email body text. No subject line. No explanations."""
    else:
        prompt = f"""You are rewriting a Job Application email to {company}.
{angle_instruction}
{project_guideline}

Acknowledge briefly that you applied earlier and are following up, but pivot the focus to this new angle.
Keep it punchy, professional, and under 120 words.
DO NOT use HTML tags like <br>. Use literal newline characters for spacing.
CRITICAL: DO NOT leave any sentence incomplete or cut off. End cleanly with a question or call to action.
CRITICAL: DO NOT include ANY sign-off, signature, "Best regards", your name, or any links/URLs at the end. The system will append the signature automatically. End the email body strictly with the final sentence's punctuation.

Resume Context:
{resume_text[:12000]}
{f"\nProject Details:\n{project_details}" if project_details else ""}

Original Email (MAY contain a signature, ignore it):
{original_body}

Return ONLY the rewritten email body text. No subject line. No explanations."""

    result_body = call_gemini(prompt, json_mode=False, temperature=0.3).strip()

    if sender_name:
        signature = f"\n\nBest regards,\n{sender_name}"
        if links_str:
            signature += f"\n{links_str}"
    elif links_str:
        signature = f"\n\n{links_str}"
    else:
        signature = ""

    return result_body + signature


def generate_follow_up(original_body: str, company: str, hr_name: str, sender_name: str = "", github_url: str = "", portfolio_url: str = "", linkedin_url: str = "") -> str:
    """Generate a polite follow-up email."""
    sender_name = sender_name.strip() if sender_name else "[Your Name]"
    links = []
    if portfolio_url and portfolio_url.strip(): links.append(f"Portfolio: {portfolio_url.strip()}")
    if github_url and github_url.strip(): links.append(f"GitHub: {github_url.strip()}")
    if linkedin_url and linkedin_url.strip(): links.append(f"LinkedIn: {linkedin_url.strip()}")
    links_str = " | ".join(links)
    prompt = f"""Write a short, polite 2-3 sentence follow-up email.
Context: I previously emailed {hr_name} at {company} about a job opportunity.
My original email: {original_body[:500]}

Return ONLY the follow-up text. No placeholders, no subject line. Sound human and genuine.
DO NOT use HTML tags like <br>. Use literal newline characters for spacing.
CRITICAL: DO NOT include ANY sign-off, signature, "Best regards", your name, or any links/URLs at the end. The system will append the signature automatically. End the email body strictly with the final sentence's punctuation."""
    result_body = call_gemini(prompt, json_mode=False, temperature=0.4).strip()

    if sender_name:
        signature = f"\n\nBest regards,\n{sender_name}"
        if links_str:
            signature += f"\n{links_str}"
    elif links_str:
        signature = f"\n\n{links_str}"
    else:
        signature = ""

    return result_body + signature

def generate_hr_reply(hr_message: str, user_instructions: str, previous_email_body: str, company: str, resume_text: str, sender_name: str = "", github_url: str = "", portfolio_url: str = "", linkedin_url: str = "") -> str:
    """Generate a reply to an HR's email based on user instructions and previous context."""
    sender_name = sender_name.strip() if sender_name else "[Your Name]"
    links = []
    if portfolio_url and portfolio_url.strip(): links.append(f"Portfolio: {portfolio_url.strip()}")
    if github_url and github_url.strip(): links.append(f"GitHub: {github_url.strip()}")
    if linkedin_url and linkedin_url.strip(): links.append(f"LinkedIn: {linkedin_url.strip()}")
    links_str = " | ".join(links)
    prompt = f"""Write a professional reply to a Hiring Manager / HR at {company}.

My Previous Email:
{previous_email_body[:2000]}

HR's Email:
{hr_message[:3000]}

My Quick Instructions / Answers for the Reply:
{user_instructions}

My Resume Context (for any additional info if needed):
{resume_text[:12000]}

Task: Draft the email reply body answering the HR's questions directly using my instructions. Keep it polite, professional, and concise.
DO NOT use HTML tags like <br>. Use literal newline characters for spacing.
CRITICAL: DO NOT include ANY sign-off, signature, "Best regards", your name, or any links/URLs at the end. The system will append the signature automatically. End the email body strictly with the final sentence's punctuation."""
    
    result_body = call_gemini(prompt, json_mode=False, temperature=0.3).strip()

    if sender_name:
        signature = f"\n\nBest regards,\n{sender_name}"
        if links_str:
            signature += f"\n{links_str}"
    elif links_str:
        signature = f"\n\n{links_str}"
    else:
        signature = ""

    return result_body + signature


def validate_email_body(body: str, sender_name: str = "") -> str:
    """Post-generation validation: strip placeholders, HTML tags, and duplicate signatures."""
    import re as _re

    # Strip leaked HTML tags
    body = _re.sub(r'<br\s*/?>', '\n', body)
    body = _re.sub(r'</?(?:p|div|span|b|i|u|strong|em)[^>]*>', '', body)

    # Replace common leftover placeholders
    placeholders = ['[Your Name]', '[your name]', '[HR Name]', '[Hiring Manager Name]',
                    '[Company Name]', '[Your Email]', '[Your Phone]']
    for ph in placeholders:
        if ph.lower() == '[your name]' and sender_name:
            body = body.replace(ph, sender_name)
            body = body.replace(ph.title(), sender_name)
        else:
            body = body.replace(ph, '')

    # Remove duplicate "Best regards" blocks (keep only the last one)
    parts = _re.split(r'(?i)\n\n\s*best\s+regards\s*,?\s*\n', body)
    if len(parts) > 2:
        # Keep everything before the first "Best regards" + the last signature block
        body = parts[0] + '\n\nBest regards,\n' + parts[-1]

    # Clean up excessive blank lines (3+ newlines -> 2)
    body = _re.sub(r'\n{3,}', '\n\n', body)

    return body.strip()
